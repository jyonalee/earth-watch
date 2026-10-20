"""Global Carbon Budget official release (XLSX)."""

from __future__ import annotations

import io
from typing import Any

from earth_watch.fetchers.base import BaseFetcher, SourceValidationError
from earth_watch.sources import GCB_GLOBAL_BUDGET_2025

_SHEET_NAME = "Global Carbon Budget"
_HEADER_YEAR = "Year"
_COL_FOSSIL = "fossil emissions excluding carbonation"
_COL_LUC = "land-use change emissions"


class GcbBudgetFetcher(BaseFetcher):
    spec = GCB_GLOBAL_BUDGET_2025

    def parse(self, content: bytes) -> list[dict[str, Any]]:
        import openpyxl

        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if _SHEET_NAME not in workbook.sheetnames:
            raise SourceValidationError(f"gcb: missing sheet {_SHEET_NAME!r}")
        worksheet = workbook[_SHEET_NAME]
        header_row_idx = None
        col_index: dict[str, int] = {}
        for idx, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if not row:
                continue
            first = row[0]
            if first == _HEADER_YEAR:
                header_row_idx = idx
                for col_i, name in enumerate(row):
                    if name:
                        col_index[str(name).strip()] = col_i
                break
        if header_row_idx is None or _COL_FOSSIL not in col_index or _COL_LUC not in col_index:
            raise SourceValidationError("gcb: could not locate budget header row")

        rows: list[dict[str, Any]] = []
        for row in worksheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
            if not row or row[0] is None:
                continue
            try:
                year = int(row[0])
            except (TypeError, ValueError):
                continue
            fossil = row[col_index[_COL_FOSSIL]]
            luc = row[col_index[_COL_LUC]]
            if fossil is None or luc is None:
                continue
            rows.append(
                {
                    "year": year,
                    "fossil_gtc": float(fossil),
                    "land_use_gtc": float(luc),
                }
            )
        return rows

    def validate(self, frame: list[dict[str, Any]]) -> None:
        if len(frame) < 50:
            raise SourceValidationError("gcb: too few budget years")
        years = [r["year"] for r in frame]
        if min(years) < 1950 or max(years) > 2100:
            raise SourceValidationError("gcb: year out of range")
        fossil = [r["fossil_gtc"] for r in frame]
        if min(fossil) < 0 or max(fossil) > 50:
            raise SourceValidationError("gcb: fossil emissions out of range")

    def compute_aggregates(self, frame: list[dict[str, Any]]) -> dict[str, Any]:
        latest = frame[-1]
        return {
            "row_count": len(frame),
            "year_min": frame[0]["year"],
            "year_max": latest["year"],
            "latest": {
                "year": latest["year"],
                "fossil_gtc": latest["fossil_gtc"],
                "land_use_gtc": latest["land_use_gtc"],
            },
            "land_use_sigma_gtc": 0.7,
            "fossil_relative_sigma": 0.05,
        }
